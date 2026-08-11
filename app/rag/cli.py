"""agent 调混合 RAG 的 CLI 工具。

用法(项目根目录下):
  python -m app.rag.cli check --title "古驰同款项链" --cost 8 --market ph
  python -m app.rag.cli copy  --market th --category earrings --title "珍珠耳环 复古"
  python -m app.rag.cli check-table "workflow/runs/xx/ph/PH_upload_top10.xlsx"
  python -m app.rag.cli learn-upload --run-dir "RoseSeek/runs/seller_upload_*_PH" \\
        --error-text "销售属性值名称不应超过 50 个字符，请检查属性值名称：..."

stdout 输出单行 JSON(agent 好解析),人类可读摘要走 stderr。
exit code:check/check-table 不通过返回 1(可作闸门)。
"""
import argparse
import json
import os
import re
import sys

from .rules import evaluate_product, supported_markets
from .vector import retrieve_copy


def main(argv):
    ap = argparse.ArgumentParser(prog="rag",
                                 description="混合RAG:规则层 check + 文案向量层 copy")
    sub = ap.add_subparsers(dest="cmd", required=True)

    c = sub.add_parser("check", help="结构化规则过滤:价格/违禁/品类/风格")
    c.add_argument("--title", required=True, help="1688 中文标题")
    c.add_argument("--cost", type=float, default=None, help="成本(人民币)")
    c.add_argument("--category", default=None, help="显式品类覆盖")
    c.add_argument("--market", default=None, help="ph/th/vn")
    c.add_argument("--min-cost", type=float, default=None)
    c.add_argument("--max-cost", type=float, default=None)

    cp = sub.add_parser("copy", help="向量检索优秀文案范例(混合 BM25+向量)")
    cp.add_argument("--market", default="ph", help="ph/th/vn")
    cp.add_argument("--category", default=None)
    cp.add_argument("--style", default=None, help="style_code 精确过滤")
    cp.add_argument("--title", default="", help="新品标题(中英皆可,作查询)")
    cp.add_argument("--top-k", type=int, default=3)

    ct = sub.add_parser("check-table", help="上传规格校验:读真上架表 xlsx 拦坏表")
    ct.add_argument("tables", nargs="+", help="上架表 xlsx 路径(可多个)")
    ct.add_argument("--max-title-dup", type=int, default=0,
                    help="重复标题告警阈值(>0 时输出 warnings;0 只做规格校验)")

    lu = sub.add_parser("learn-upload", help="上传失败反学:记 history + 升级规则")
    lu.add_argument("--run-dir", default=None, help="上传 run 目录(含 seller_upload_events.jsonl)")
    lu.add_argument("--error-text", default=None,
                    help="TikTok 真实失败原因(从解析错误页读,逐件原因靠它)")
    lu.add_argument("--market", default=None, help="市场(缺省从 run-dir 名推断)")

    args = ap.parse_args(argv)

    if args.cmd == "check":
        r = evaluate_product(args.title, cost_cny=args.cost, market=args.market,
                             min_cost=args.min_cost, max_cost=args.max_cost)
        print(json.dumps(r, ensure_ascii=False, separators=(",", ":")))
        sys.stderr.write("规则判定:%s | %s\n"
                         % ("通过" if r["passed"] else "不通过", r["summary"]))
        return 0 if r["passed"] else 1

    if args.cmd == "check-table":
        return _check_table(args.tables, args.max_title_dup)

    if args.cmd == "learn-upload":
        return _learn_upload(args.run_dir, args.error_text, args.market)

    ex = retrieve_copy(args.market, args.category, args.style, args.title, args.top_k)
    print(json.dumps({"market": (args.market or "").upper(),
                      "category": args.category,
                      "count": len(ex),
                      "examples": ex},
                     ensure_ascii=False, separators=(",", ":")))
    sys.stderr.write("检索 %d 条文案范例(market=%s)\n"
                     % (len(ex), (args.market or "").upper()))
    return 0


# ---------------------------------------------------------------- check-table
def _check_table(tables, max_title_dup=0):
    """读真上架表,按 upload_learning 反学规则校验上传规格,上架前拦坏表。

    扫描逻辑与 app/agent/approval._count_products 一致:Template sheet 里
    找含 product_name 的英文表头行,price 列可解析成数字才算真商品行。
    规格(规则可配):product_name≤255 / property_name_*≤20 / property_value_*≤50。
    """
    from .rules import check_upload_spec
    import openpyxl

    fails, warnings = [], []
    total = 0
    for path in tables:
        if not os.path.exists(path):
            fails.append({"table": path, "row": 0, "column": "(file)",
                          "value": path, "length": 0, "max_len": 0,
                          "rule_id": "missing_file", "tiktok_error": "上架表不存在"})
            continue
        try:
            wb = openpyxl.load_workbook(path, read_only=True)
        except Exception as e:
            fails.append({"table": path, "row": 0, "column": "(file)",
                          "value": str(e), "length": 0, "max_len": 0,
                          "rule_id": "unreadable", "tiktok_error": "表无法读取"})
            continue
        ws = wb["Template"] if "Template" in wb.sheetnames else wb.active
        rows = ws.iter_rows(values_only=True)
        header = None
        for r in rows:
            if any(isinstance(c, str) and c.strip() == "product_name" for c in (r or [])):
                header = [str(c) if c is not None else "" for c in r]
                break
        if not header:
            wb.close()
            fails.append({"table": path, "row": 0, "column": "(sheet)",
                          "value": "no product_name header", "length": 0, "max_len": 0,
                          "rule_id": "no_header", "tiktok_error": "找不到 product_name 表头"})
            continue
        idx = {h.strip(): i for i, h in enumerate(header) if h.strip()}
        pi = idx.get("product_name")
        pr = idx.get("price")
        skip = {"create_product", "metric", "category_v2", "商品名称", "产品名称",
                "必填", "选填", "商品描述", "产品描述"}
        seen_titles = set()
        # 模板 property 列不止 1/2,动态扫全部 property_name_/property_value_*
        prop_name_cols = sorted(i for h, i in idx.items() if h.startswith("property_name_"))
        prop_value_cols = sorted(i for h, i in idx.items() if h.startswith("property_value_"))
        for r in rows:
            if not r or not any(c is not None and str(c).strip() for c in r):
                continue
            title = str(r[pi]).strip() if (pi is not None and pi < len(r) and r[pi] is not None) else ""
            if not title or title in skip:
                continue
            price = r[pr] if (pr is not None and pr < len(r)) else None
            try:
                float(price)
            except (TypeError, ValueError):
                continue
            total += 1
            row_label = title[:40]
            checks = [("product_name", title)]
            checks += [("property_name", str(r[i]) if i < len(r) and r[i] is not None else "")
                       for i in prop_name_cols]
            checks += [("property_value", str(r[i]) if i < len(r) and r[i] is not None else "")
                       for i in prop_value_cols]
            for col, val in checks:
                if not val:
                    continue
                res = check_upload_spec(val, col)
                if not res.get("passed"):
                    fails.append({"table": path, "row": total, "product_name": row_label,
                                  "column": col, "value": val[:120],
                                  "length": len(val), "max_len": res.get("max_len"),
                                  "rule_id": res.get("rule_id"),
                                  "tiktok_error": res.get("tiktok_error")})
            if max_title_dup > 0 and title in seen_titles:
                warnings.append({"table": path, "product_name": row_label,
                                 "detail": "重复标题(导出层已按 title_cn 去重,此处仅提示)"})
            seen_titles.add(title)
        wb.close()

    r = {"stage": "check_table", "passed": not fails, "rows": total,
         "fails": fails, "warnings": warnings}
    print(json.dumps(r, ensure_ascii=False, separators=(",", ":")))
    sys.stderr.write("表规校验:%d 行, %d FAIL%s\n"
                     % (total, len(fails),
                        ", %d 提示" % len(warnings) if warnings else ""))
    return 0 if not fails else 1


# ---------------------------------------------------------------- learn-upload
def _upload_learning_path():
    from ..config import CONFIG_DIR
    return os.path.join(str(CONFIG_DIR), "upload_learning.json")


def _load_upload_learning(path):
    from .io import read_json
    if os.path.exists(path):
        return read_json(path) or {}
    return {"rules": [], "history": []}


def _save_upload_learning(path, data):
    import tempfile
    d = os.path.dirname(path) or "."
    fd, tmp = tempfile.mkstemp(dir=d, suffix=".json")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            f.write("\n")
        os.replace(tmp, path)
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise


def _learn_upload(run_dir=None, error_text=None, market=None):
    """上传失败反学:读事件文件记 history,按 --error-text 确认/新增规则。

    诚实边界:事件文件只记计数(error_products/ready_products),逐件原因
    靠 --error-text(从 TikTok 解析错误页读)传入;脚本不编造原因。
    """
    import datetime
    path = _upload_learning_path()
    data = _load_upload_learning(path)
    rules = data.setdefault("rules", [])
    history = data.setdefault("history", [])
    learned = []

    if run_dir:
        events_path = os.path.join(run_dir, "seller_upload_events.jsonl")
        page_path = os.path.join(run_dir, "seller_upload_page.json")
        base = os.path.basename(os.path.normpath(run_dir))
        mkt = market or ""
        if not mkt:
            for part in base.split("_"):
                if part.upper() in ("PH", "TH", "VN"):
                    mkt = part.upper()
                    break
        parse_err = None
        page_text = ""
        if os.path.exists(events_path):
            for line in open(events_path, encoding="utf-8-sig"):
                line = line.strip()
                if not line:
                    continue
                try:
                    ev = json.loads(line)
                except Exception:
                    continue
                if (ev.get("phase") == "parse" and ev.get("status") == "complete"
                        and int((ev.get("data") or {}).get("error_products") or 0) > 0):
                    parse_err = ev
        if os.path.exists(page_path):
            try:
                page_text = str(json.load(open(page_path, encoding="utf-8-sig")))[:500]
            except Exception:
                page_text = ""
        if parse_err:
            d = parse_err.get("data") or {}
            entry = {
                "run": d.get("upload_file", "?"),
                "market": mkt or "?",
                "time": parse_err.get("timestamp", "?"),
                "error_products": int(d.get("error_products") or 0),
                "ready_products": int(d.get("ready_products") or 0),
                "page": page_text,
                "note": "learn-upload 自动记录(计数;逐件原因需 --error-text)",
            }
            key = (entry["time"], entry["error_products"])
            if not any((h.get("time"), h.get("error_products")) == key for h in history):
                history.append(entry)
                learned.append("history:" + entry["run"])

    if error_text and error_text.strip():
        text = error_text.strip()
        nums = re.findall(r"\d+", text)
        matched = None
        for rule in rules:
            err = rule.get("tiktok_error") or ""
            if err and err in text:
                matched = rule.get("id")
                break
        if matched is None:
            max_len = int(nums[0]) if nums else 0
            for rule in rules:
                if int(rule.get("max_len") or 0) == max_len:
                    matched = rule.get("id")
                    break
        if matched:
            learned.append("matched_rule:" + matched)
            sys.stderr.write("失败原因命中已有规则 %s\n" % matched)
        elif not any(r.get("tiktok_error") == text for r in rules):
            # 未匹配 → 新候选规则(severity=manual 不自动拦截,人工确认后生效)
            rules.append({
                "id": "learned_%d" % (len(rules) + 1),
                "kind": "tiktok_error", "column": "",
                "max_len": int(nums[0]) if nums else 0,
                "tiktok_error": text,
                "source": "learn-upload --error-text",
                "learned_at": datetime.date.today().isoformat(),
                "severity": "manual",
            })
            learned.append("new_rule:" + text[:40])

    _save_upload_learning(path, data)
    r = {"stage": "learn_upload", "path": path, "learned": learned,
         "rules": len(rules), "history": len(history)}
    print(json.dumps(r, ensure_ascii=False, separators=(",", ":")))
    sys.stderr.write("反学:%s(规则 %d 条,历史 %d 条)\n"
                     % ("; ".join(learned) or "无新增", len(rules), len(history)))
    return 0


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(main(sys.argv[1:]))
