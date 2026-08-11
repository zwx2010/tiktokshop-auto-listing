"""Excel 导出 —— 看板「导出Excel」按钮与 tools/export_products_excel.py 共用。

只做一件事：把数据库商品/SKU 渲染成 openpyxl Workbook。
数据读取方（API 或脚本）负责传入 Session。
"""
import csv
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import selectinload

from .config import BASE_DIR
from .models import Listing, Product
from .pricing import price_skus
from .timeutil import fmt_utc

HEADER_FILL = PatternFill("solid", fgColor="2F5597")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)


def _codex_goods_ids() -> set[str]:
    """工作流 PH codex_listing_copy.csv 的 goods_id 集合（选品表标注是否已有地道文案）。

    文件缺失（工作流未生成过文案）时返回空集，所有商品标「无」。跨项目路径失败不抛异常。
    """
    p = BASE_DIR.parent / "tk自动化工作流" / "runs" / "smoke_01" / "ph" / "codex_listing_copy.csv"
    if not p.exists():
        return set()
    try:
        with open(p, encoding="utf-8-sig") as f:
            return {str(r["goods_id"]).strip() for r in csv.DictReader(f) if r.get("goods_id")}
    except Exception:
        return set()


def _style_header(ws, headers, widths):
    for col, (name, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    ws.row_dimensions[1].height = 22


def build_workbook(db) -> Workbook:
    """把全部商品/上架情况渲染成双 sheet 的 Workbook。"""
    products = (
        db.query(Product)
        .options(selectinload(Product.skus))
        .order_by(Product.id.desc())
        .all()
    )
    # 每个商品上架过的店铺数（去重）
    listing_counts: dict[int, set] = {}
    for pid, acc_id in db.query(Listing.product_id, Listing.account_id).all():
        listing_counts.setdefault(pid, set()).add(acc_id)

    wb = Workbook()

    # ---------- Sheet1 商品（兼作选品专用表：含 PH 定价预览 / SKU 数 / Codex 文案标记） ----------
    ws = wb.active
    ws.title = "商品"
    codex_gids = _codex_goods_ids()
    headers = [
        "ID", "商品ID(goods_id)", "标题", "分类", "成本(¥)", "成本来源",
        "图片数", "主图URL", "已上架店铺数", "状态", "创建时间",
        "PH展示价", "PH目标成交价", "SKU数", "Codex文案",
    ]
    widths = [6, 15, 42, 12, 9, 14, 8, 60, 12, 10, 19, 12, 12, 8, 10]
    _style_header(ws, headers, widths)

    for r, p in enumerate(products, start=2):
        ws.cell(r, 1, p.id)
        ws.cell(r, 2, p.source_goods_id)
        ws.cell(r, 3, p.title_cn)
        ws.cell(r, 4, p.category)
        ws.cell(r, 5, p.cost_cny_used)
        ws.cell(r, 6, p.cost_source)
        ws.cell(r, 7, len(p.image_urls) if p.image_urls else 0)

        # 选品用：PH 定价预览（代表成本）+ SKU 数 + 是否已有 Codex 文案
        pricing = None
        if p.cost_cny_used > 0:
            try:
                pricing = price_skus("PH", p.cost_cny_used, p.weight_g or 80)
            except Exception:
                pricing = None
        ws.cell(r, 12, pricing["display_price"] if pricing else "")
        ws.cell(r, 13, pricing["target_sale_price"] if pricing else "")
        ws.cell(r, 14, len(p.skus))
        ws.cell(r, 15, "有" if p.source_goods_id in codex_gids else "无")

        main_img = p.main_image_url or (p.image_urls[0] if p.image_urls else "")
        cell = ws.cell(r, 8, main_img)
        if main_img.startswith("http"):
            cell.hyperlink = main_img
            cell.style = "Hyperlink"

        ws.cell(r, 9, len(listing_counts.get(p.id, set())))
        ws.cell(r, 10, p.status)
        ws.cell(r, 11, fmt_utc(p.created_at, "%Y-%m-%d %H:%M"))

        # 点标题跳到拼多多源页
        if p.source_goods_id:
            url = f"https://mobile.pinduoduo.com/goods.html?goods_id={p.source_goods_id}"
            ws.cell(r, 3).hyperlink = url
            ws.cell(r, 3).style = "Hyperlink"

    # ---------- Sheet2 SKU明细 ----------
    ws2 = wb.create_sheet("SKU明细")
    headers2 = ["商品ID", "标题", "颜色", "规格", "尺寸", "供应商SKU", "成本(¥)", "库存", "采集时间"]
    widths2 = [10, 42, 12, 12, 10, 16, 9, 8, 19]
    _style_header(ws2, headers2, widths2)

    r2 = 2
    for p in products:
        created_str = fmt_utc(p.created_at, "%Y-%m-%d %H:%M")
        if not p.skus:
            ws2.cell(r2, 1, p.id)
            ws2.cell(r2, 2, p.title_cn)
            ws2.cell(r2, 7, p.cost_cny_used)
            ws2.cell(r2, 9, created_str)
            r2 += 1
            continue
        for s in p.skus:
            ws2.cell(r2, 1, p.id)
            ws2.cell(r2, 2, p.title_cn)
            ws2.cell(r2, 3, s.color)
            ws2.cell(r2, 4, s.style)
            ws2.cell(r2, 5, s.size)
            ws2.cell(r2, 6, s.supplier_sku_id)
            ws2.cell(r2, 7, s.cost_cny)
            ws2.cell(r2, 8, s.stock)
            ws2.cell(r2, 9, created_str)
            r2 += 1

    return wb


def workbook_bytes(db) -> bytes:
    """序列化 Workbook 为 xlsx 字节流（API 下载用）。"""
    buf = BytesIO()
    build_workbook(db).save(buf)
    buf.seek(0)
    return buf.getvalue()
