#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""从英文版 TikTok 模板提取官方英文类目路径,并按 category_id 与中文模板对齐生成映射。

背景:上架表(上传 xlsx)不能含汉字。中文模板的 Category 下拉是中文路径
(如「平价饰品/耳环」),官方英文模板的 Category 下拉是英文路径。两个模板的
Category sheet 都是「路径 | category_id」两列,而 category_id 是语言无关的权威键
(Example sheet 里类目单元格存的就是「路径 (id)」格式)。因此:
  1. 读中文模板 Category sheet → {id: 中文路径}
  2. 读英文模板 Category sheet → {id: 英文路径}
  3. 按 id 对齐 → {中文路径: 英文路径}(官方值,逐字匹配英文模板下拉)

产出 {工作流}/config/en_category_map.json,供 Fill-TikTokTemplate.ps1 的 -English
模式把中文解析结果换成官方英文路径。ps1 在映射缺失/查不到时硬报错,绝不出错表。

用法(平台根目录下):
    python scripts/extract_en_category_map.py --market ph
    python scripts/extract_en_category_map.py            # 全部市场(ph/th/vn)
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

PLATFORM_ROOT = Path(__file__).resolve().parents[1]
WF_ROOT = Path(__file__).resolve().parents[2] / "tk自动化工作流"


def load_category_sheet(template: Path) -> dict[int, str]:
    """Category sheet → {category_id(int): 路径}。路径与 id 都必须非空。

    必须 read_only=False:英文模板的 Category sheet 在 read_only 模式下
    因 dimension 属性异常只读出 1 行,普通模式才能读全 59 行。
    """
    wb = openpyxl.load_workbook(template, read_only=False, data_only=True)
    if "Category" not in wb.sheetnames:
        sys.exit(f"模板缺 Category sheet: {template}")
    ws = wb["Category"]
    out: dict[int, str] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or row[1] is None:
            continue
        path = str(row[0]).strip()
        try:
            cid = int(float(row[1]))
        except (TypeError, ValueError):
            continue
        if path and cid > 0:
            out[cid] = path
    return out


def extract_en_genders(template: Path) -> list[str]:
    """英文模板 HiddenAttr 列 J 的去重非空值(性别选项)。列 I 为类目分组键。"""
    wb = openpyxl.load_workbook(template, read_only=False, data_only=True)
    if "HiddenAttr" not in wb.sheetnames:
        return []
    ws = wb["HiddenAttr"]
    vals: list[str] = []
    seen: set[str] = set()
    for row in ws.iter_rows(values_only=True):
        if row is None or len(row) < 10:
            continue
        v = row[9]  # J 列(0 基第 10 列)
        if v is None:
            continue
        s = str(v).strip()
        if s and s not in seen:
            seen.add(s)
            vals.append(s)
    return vals


def main() -> int:
    ap = argparse.ArgumentParser(description="英文模板类目映射(按 category_id 对齐中文模板)")
    ap.add_argument("--market", default="", choices=["", "ph", "th", "vn"], help="市场,留空=全部")
    args = ap.parse_args()

    markets = [args.market] if args.market else ["ph", "th", "vn"]
    output: dict[str, dict[str, str]] = {}
    genders: dict[str, list[str]] = {}
    any_en_missing = False

    for mkt in markets:
        cn_tpl = WF_ROOT / "templates" / f"accessories_{mkt}.xlsx"
        en_tpl = WF_ROOT / "templates" / f"accessories_{mkt}_en.xlsx"
        if not cn_tpl.exists():
            print(f"[跳过] {mkt}: 缺中文模板 {cn_tpl.name}")
            continue
        if not en_tpl.exists():
            any_en_missing = True
            print(f"[未下载] {mkt}: 缺英文模板 {en_tpl.name} —— 请到卖家中心切英文界面后下载到 templates/")
            continue

        cn = load_category_sheet(cn_tpl)
        en = load_category_sheet(en_tpl)
        # 对齐校验:id 应一一对应
        cn_only = [i for i in cn if i not in en]
        en_only = [i for i in en if i not in cn]
        if cn_only:
            print(f"[警告] {mkt}: 中文模板有、英文模板无的 category_id: {cn_only}")
        if en_only:
            print(f"[警告] {mkt}: 英文模板有、中文模板无的 category_id: {en_only}")

        mapping = {}
        for cid, cn_path in cn.items():
            en_path = en.get(cid)
            if not en_path:
                continue
            if cn_path == en_path:
                continue  # 本身就是英文的,无需映射(理论上不会有)
            mapping[cn_path] = en_path
        output[mkt] = mapping
        genders[mkt] = extract_en_genders(en_tpl)
        print(f"[OK] {mkt}: {len(mapping)} 条 中文路径→官方英文路径(按 id 对齐)")
        # 抽样核对几条常见配饰类目
        for sample in ("平价饰品/耳环", "平价饰品/项链", "发饰/头绳"):
            if sample in mapping:
                print(f"      {sample}  ->  {mapping[sample]}")

    if any_en_missing:
        print("\n英文模板尚未就位,本次未产出完整映射。下载后重跑本脚本即可。")
    if not output:
        print("没有任何市场可用,未写入映射文件。")
        return 1

    out_path = WF_ROOT / "config" / "en_category_map.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(
        json.dumps({"categories": output, "genders": genders}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"\n写入映射: {out_path}")
    return 0 if not any_en_missing else 2


if __name__ == "__main__":
    sys.exit(main())
